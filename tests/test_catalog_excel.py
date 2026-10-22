"""Excel -> components.json converter. Uses a copy of the real sheet, no network."""
import shutil

from openpyxl import load_workbook

from core.components import DATA_FILE, load_components
from tools.catalog_from_excel import DEFAULT_XLSX, read_sheet


def _sheet(tmp_path, **fill):
    """Copy the real sheet and fill one part row: fill={'LM317T': {5: 'C1', ...}} (column -> value)."""
    path = tmp_path / "parts.xlsx"
    shutil.copy(DEFAULT_XLSX, path)
    wb = load_workbook(path)
    ws = wb["Parts"]
    rows = {ws.cell(row=r, column=1).value: r for r in range(2, 40)}
    for key, cells in fill.items():
        for col, value in cells.items():
            ws.cell(row=rows[key], column=col, value=value)
    wb.save(path)
    return path


def test_untouched_sheet_matches_current_catalog(tmp_path):
    entries, problems, todo = read_sheet(_sheet(tmp_path))
    assert problems == []
    assert {e["mpn"].upper() for e in entries} == set(load_components())

def test_filled_row_becomes_an_entry(tmp_path):
    path = _sheet(tmp_path, LM317T={5: "C12345", 6: "TO-220", 7: "1, 2, 3", 8: 40, 9: 1.5})
    entries, problems, _ = read_sheet(path)
    lm = next(e for e in entries if e["mpn"] == "LM317T")
    assert problems == []
    assert lm["lcsc"] == "C12345" and lm["pins"] == {"1": "ADJ", "2": "OUT", "3": "IN"}


def test_mistakes_are_explained_with_the_row(tmp_path):
    path = _sheet(tmp_path, NE555P={5: "12345", 6: "SOIC-8", 7: "1,2,3", 8: 16})
    _, problems, _ = read_sheet(path)
    text = "\n".join(problems)
    assert "NE555P" in text and "C12345" in text and "surface-mount" in text and "3 pin numbers" in text


def test_resistor_values_fill_lcsc_by_value(tmp_path):
    path = _sheet(tmp_path)
    wb = load_workbook(path)
    wb["Resistors"]["C3"] = "C22222"            # row 3 = RES-0.25W 240
    wb.save(path)
    entries, problems, _ = read_sheet(path)
    res = next(e for e in entries if e["mpn"] == "RES-0.25W")
    assert problems == [] and res["lcsc_by_value"] == {"240": "C22222"} and res["power_w"] == 0.25


def test_components_file_still_valid():
    assert DATA_FILE.exists()
