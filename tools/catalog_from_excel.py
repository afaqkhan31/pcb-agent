"""Turn the Excel parts sheet into data/components.json.

Run:   python -m tools.catalog_from_excel
       python -m tools.catalog_from_excel path/to/other.xlsx

The Excel file (data/catalog_input.xlsx) is the ONLY place the catalog is edited.
This script checks every row, explains problems in plain words (with the Excel row
number), and only writes components.json when everything filled in is valid.
Owner: Hamna (catalog).  Hamdan reviews pin names.
"""
from __future__ import annotations

import json
import re
import sys
from pathlib import Path

from openpyxl import load_workbook

from core.components import DATA_FILE, LCSC_PATTERN

DEFAULT_XLSX = DATA_FILE.parent / "catalog_input.xlsx"
SMD_WORDS = re.compile(r"SMD|0402|0603|0805|1206|SOT|SOD|SOIC|SOP|QFN|SMA\b|SMB\b|SMC\b", re.I)
GENERIC_POWER_W = {"RES-0.25W": 0.25, "RES-1W": 1.0}     # resistor families: one entry + values

PARTS_HEADERS = ["Part key", "Category", "Pin names", "Typical pin numbers", "LCSC number", "Package",
                 "Pin numbers", "Max voltage (V)", "Max current (A)", "Notes", "Status", "Description"]


def _text(value) -> str:
    return "" if value is None else str(value).strip()


def _number(value, where: str, problems: list[str]):
    if value in (None, ""):
        return None
    try:
        return float(value) if float(value) != int(float(value)) else int(float(value))
    except (TypeError, ValueError):
        problems.append(f"{where}: '{value}' is not a number.")
        return None


def _split(cell: str) -> list[str]:
    return [x.strip() for x in cell.replace(";", ",").split(",") if x.strip()]


def read_sheet(path: Path) -> tuple[list[dict], list[str], int]:
    """Return (catalog entries, problems, number of rows still to do)."""
    wb = load_workbook(path, data_only=True)
    problems: list[str] = []

    # --- resistor values (second sheet) -> {family: {value: C-number}}
    by_value: dict[str, dict[str, str]] = {}
    if "Resistors" in wb.sheetnames:
        for n, row in enumerate(wb["Resistors"].iter_rows(min_row=2, values_only=True), start=2):
            family, value, lcsc = _text(row[0]), _text(row[1]), _text(row[2]).upper()
            if not (family or value or lcsc):
                continue
            if not lcsc:
                continue                      # value listed but no C-number yet: still to do
            if not (family and value):
                problems.append(f"Resistors row {n}: has an LCSC number but no family or value.")
                continue
            if not LCSC_PATTERN.match(lcsc):
                problems.append(f"Resistors row {n}: '{lcsc}' is not an LCSC number (should look like C12345).")
                continue
            by_value.setdefault(family, {})[value] = lcsc

    entries, todo, seen = [], 0, set()
    ws = wb["Parts"]
    header = [_text(c.value) for c in ws[1]]
    col = {name: header.index(name) for name in PARTS_HEADERS if name in header}
    missing_cols = [h for h in PARTS_HEADERS if h not in col]
    if missing_cols:
        return [], [f"The Parts sheet is missing columns: {missing_cols}. Don't rename or delete headers."], 0

    for n, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        get = lambda name: row[col[name]] if col[name] < len(row) else None   # noqa: E731
        key = _text(get("Part key"))
        if not key:
            continue
        where = f"Parts row {n} ({key})"
        if key.upper() in seen:
            problems.append(f"{where}: this part key appears twice.")
            continue
        seen.add(key.upper())

        names = _split(_text(get("Pin names")))
        lcsc = _text(get("LCSC number")).upper()
        package = _text(get("Package"))
        numbers = _split(_text(get("Pin numbers")))
        generic = key in GENERIC_POWER_W

        # Nothing filled in yet -> still to do, not an error.
        if generic and not by_value.get(key):
            todo += 1
            continue
        if not generic and not any([lcsc, package, numbers]):
            todo += 1
            continue

        before = len(problems)
        if not generic:
            if not LCSC_PATTERN.match(lcsc):
                problems.append(f"{where}: LCSC number '{lcsc}' should look like C12345.")
            if not numbers:
                problems.append(f"{where}: fill in the pin numbers from the datasheet.")
        if not package:
            problems.append(f"{where}: fill in the package (for example TO-220, DO-41, DIP-8).")
        elif SMD_WORDS.search(package):
            problems.append(f"{where}: package '{package}' looks surface-mount. We only use through-hole parts.")
        if generic:
            numbers = numbers or ["1", "2"]
        if numbers and len(numbers) != len(names):
            problems.append(f"{where}: {len(names)} pin names ({', '.join(names)}) but "
                            f"{len(numbers)} pin numbers. They must match one to one, in the same order.")
        if len(set(numbers)) != len(numbers):
            problems.append(f"{where}: a pin number is used twice.")
        v_max = _number(get("Max voltage (V)"), where, problems)
        i_max = _number(get("Max current (A)"), where, problems)
        if v_max is None and not generic:
            problems.append(f"{where}: fill in the maximum voltage from the datasheet.")
        if len(problems) > before:
            continue

        entry = {
            "mpn": key, "lcsc": None if generic else lcsc,
            "category": _text(get("Category")) or "other",
            "description": _text(get("Description")) or key,
            "package": package,
            "pins": dict(zip(numbers, names)),
            "v_max": v_max if v_max is not None else 250,
            "i_max_a": i_max,
            "power_w": GENERIC_POWER_W.get(key),
            "notes": _text(get("Notes")),
        }
        if generic:
            entry["lcsc_by_value"] = by_value[key]
        entries.append(entry)
    return entries, problems, todo


def main(argv: list[str]) -> int:
    path = Path(argv[1]) if len(argv) > 1 else DEFAULT_XLSX
    if not path.exists():
        print(f"Cannot find {path}")
        return 1
    entries, problems, todo = read_sheet(path)
    if problems:
        print("Fix these in Excel, save, and run again (nothing was changed):\n")
        for p in problems:
            print("  -", p)
        return 1
    DATA_FILE.write_text(json.dumps(entries, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    print(f"Wrote {len(entries)} parts to {DATA_FILE.name} ({todo} rows still empty in Excel).\n")
    from core.catalog_todo import NEEDED, todo as missing
    lines = missing()
    print(f"Template parts ready: {len(NEEDED) - len({l.split()[1] for l in lines})}/{len(NEEDED)}")
    print("\n".join(lines) if lines else "All template parts are in the catalog.")
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv))
